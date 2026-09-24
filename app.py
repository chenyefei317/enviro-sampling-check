import datetime
import io
import os
import re
import zipfile
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor, Pt
import pandas as pd
import qrcode
from PIL import Image
import requests
import streamlit as st
from streamlit_drawable_canvas import st_canvas

# ================= 1. 页面配置与初始化 =================
st.set_page_config(
    page_title="自行监测与采样人员现场点检平台",
    layout="centered",
    initial_sidebar_state="expanded",
)

# 注入华文宋体全局样式与隐藏默认元素
hide_streamlit_style = """
    <style>
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}
    html, body, [class*="css"] {
        font-family: "华文宋体", SimSun, serif;
    }
    </style>
    """
st.markdown(hide_streamlit_style, unsafe_allow_html=True)

# ================= 2. 动态加载 Excel 点检表要求 =================
excel_filename = "环境监测点检表 2026-06-15.xlsx"
sheet_names = []
excel_data_dict = {}

if os.path.exists(excel_filename):
  try:
    xls = pd.ExcelFile(excel_filename)
    sheet_names = xls.sheet_names
    for sheet in sheet_names:
      df = pd.read_excel(xls, sheet_name=sheet)
      items = []
      for col in df.columns:
        for val in df[col].dropna().astype(str):
          val_clean = val.strip()
          if (
              len(val_clean) > 1
              and val_clean not in items
              and not val_clean.startswith("Unnamed")
          ):
            items.append(val_clean)
      excel_data_dict[sheet] = items if items else [f"执行{sheet}相关标准点检"]
  except Exception as e:
    pass

# 若未找到 Excel，提供默认兜底选项
if not sheet_names:
  sheet_names = [
      "VOC挥发性有机物（有组织废气）",
      "颗粒物（有组织废气）",
      "生活污水",
      "雨水",
      "噪声",
  ]
  for s in sheet_names:
    excel_data_dict[s] = [
        f"采样/监测仪器在检定有效期内且完成现场校准",
        f"现场作业人员按规范正确佩戴个人防护用品",
        f"点位设置、频次及方法严格执行国家标准及方案要求",
    ]


# ================= 3. 百度网盘自动上传函数（带 OAuth2 自动刷新） =================
def refresh_baidu_access_token():
  try:
    client_id = st.secrets.get("BAIDU_CLIENT_ID", "")
    client_secret = st.secrets.get("BAIDU_CLIENT_SECRET", "")
    refresh_token = st.secrets.get("BAIDU_REFRESH_TOKEN", "")
    if not client_id or not client_secret or not refresh_token:
      return None
    token_url = "https://pan.baidu.com/oauth/2.0/token"
    params = {
        "grant_type": "refresh_token",
        "refresh_token": refresh_token,
        "client_id": client_id,
        "client_secret": client_secret,
    }
    response = requests.get(token_url, params=params)
    res_data = response.json()
    return res_data.get("access_token")
  except Exception:
    return None


def upload_to_baidu_netdisk_with_auto_refresh(file_bytes, remote_filename):
  access_token = st.secrets.get("BAIDU_ACCESS_TOKEN", "")
  if not access_token:
    return (
        False,
        "未配置网盘凭证，文件已在本地生成并可通过网页下载/ZIP打包保存。",
    )

  sub_folder = "自行监测点检档案"
  target_path = f"/apps/慧瑞EHS合规档案/{sub_folder}/{remote_filename}"

  def send_upload_request(token):
    upload_url = f"https://pan.baidu.com/rest/2.0/xpan/file?method=upload&access_token={token}&path={target_path}&uploadid=&file=1"
    files = {"file": (remote_filename, file_bytes)}
    return requests.post(upload_url, files=files).json()

  result = send_upload_request(access_token)
  if "errno" in result and result["errno"] in [110, 111]:
    new_token = refresh_baidu_access_token()
    if new_token:
      result = send_upload_request(new_token)
    else:
      return False, "Token 已过期且自动刷新失败。"

  if "errno" in result and result["errno"] == 0:
    return True, f"成功同步至网盘：/apps/慧瑞EHS合规档案/{sub_folder}/"
  else:
    return False, f"网盘上传失败: {result.get('error_msg', '未知错误')}"


# ================= 4. 侧边栏：Logo与微信分享 =================
with st.sidebar:
  try:
    st.image("logo.png", width=160)
  except Exception:
    st.image(
        "https://upload.wikimedia.org/wikipedia/commons/thumb/c/c5/Ikea_logo.svg/800px-Ikea_logo.svg.png",
        width=160,
    )

  st.markdown("### 📱 微信扫码与分享")
  st.write("已自动关联您的云端网址，二维码将实时更新供手机扫码填报。")

  app_url = st.text_input(
      "应用公网链接 (URL)", value="https://enviro-sampling-check.streamlit.app/"
  )

  if app_url:
    qr = qrcode.make(app_url)
    img_buffer = io.BytesIO()
    qr.save(img_buffer, format="PNG")
    st.image(
        Image.open(img_buffer), caption="微信扫码快速填报点检表", width=160
    )
    st.info(
        "💡 **提示**：将上方链接复制并发送至工作群，采样人员即可手机端随时完成现场自检与点报。"
    )

  st.markdown("---")
  st.markdown("### 📥 原始点检表模板下载")
  if os.path.exists(excel_filename):
    with open(excel_filename, "rb") as fe:
      st.download_button(
          label="📄 下载环境监测点检表.xlsx",
          data=fe.read(),
          file_name=excel_filename,
          use_container_width=True,
      )

# ================= 5. 主界面逻辑 =================
col_logo, col_title = st.columns([1, 6])
with col_logo:
  try:
    st.image("logo.png", width=110)
  except Exception:
    st.image(
        "https://upload.wikimedia.org/wikipedia/commons/thumb/c/c5/Ikea_logo.svg/800px-Ikea_logo.svg.png",
        width=110,
    )
with col_title:
  st.markdown("## 自行监测与采样人员现场点检平台")

st.markdown(
    "本系统用于环境及职业卫生自行监测过程中，采样人员的现场操作点检与自检合规确认。"
)

# 基础信息录入
st.subheader("1. 采样人员与任务基本信息")
col1, col2 = st.columns(2)
with col1:
  sampling_person = st.text_input("采样人员姓名 (必填)：")
with col2:
  employee_id = st.text_input(
      "身份证号 / 工号 (必填)：", help="请输入18位身份证号或有效工号"
  )

# 中文“选择全部”控制多选任务
select_all_tasks = st.checkbox("☑ 选择全部监测任务")
default_selected = (
    sheet_names if select_all_tasks else ([sheet_names[0]] if sheet_names else [])
)

selected_tasks = st.multiselect(
    "监测任务/点位 (可复选)：",
    options=sheet_names,
    default=default_selected,
    help="勾选本次需要执行监测和点检的任务类型",
)

# 若选择了噪声，特别展示国家标准指引
if any("噪声" in t for t in selected_tasks):
  with st.expander("🔊 【合规指引】噪声监测国家标准参考", expanded=True):
    st.info("""
        * **标准依据**：执行《工业企业厂界环境噪声排放标准》（**GB 12348-2008**）及《声环境质量标准》（**GB 3096-2008**）。
        * **测点要求**：厂界噪声测点应选择在工业企业厂界经受噪声影响最具代表性的位置，传声器应加防风罩，避开强磁场和振动干扰。
        * **气象条件**：测量应在无雨、无雪、风速小于 5m/s 的气象条件下进行。
        """)

st.write("---")
st.markdown("### 📋 动态检查要求与点检确认")

# 根据所选任务动态聚合 Excel 中的检查要求，避免名称重复
current_dynamic_items = []
for task in selected_tasks:
  if task in excel_data_dict:
    for req in excel_data_dict[task]:
      if task in req:
        formatted_req = f"【{req}】" if not req.startswith("【") else req
      else:
        formatted_req = f"【{task}】{req}"

      if formatted_req not in current_dynamic_items:
        current_dynamic_items.append(formatted_req)

if not current_dynamic_items:
  current_dynamic_items = [
      "【通用要求】现场采样仪器在检定有效期内且完成现场校准",
      "【通用要求】采样人员按规范正确佩戴个人防护用品",
  ]


# 稳健的“选择全部检查项”状态管理逻辑
def toggle_all_checks():
  target_state = st.session_state.get("select_all_checks_toggle", False)
  for i in range(len(current_dynamic_items)):
    st.session_state[f"chk_dyn_{i}"] = target_state


st.checkbox(
    "☑ 选择全部检查项",
    key="select_all_checks_toggle",
    on_change=toggle_all_checks,
)

user_checks = []
st.info("请对照以下对应任务的 Excel 检查要求逐项核对：")
for idx, req_text in enumerate(current_dynamic_items):
  # 初始化每个复选框的 Session State
  if f"chk_dyn_{idx}" not in st.session_state:
    st.session_state[f"chk_dyn_{idx}"] = False

  isChecked = st.checkbox(req_text, key=f"chk_dyn_{idx}")
  user_checks.append(isChecked)

# ================= 6. 手写签名与手写日期栏（并排双画布） =================
current_date_str = datetime.date.today().strftime("%Y年%m月%d日")

st.write("---")
st.subheader("✍️ 2. 采样人员手写签名与手写日期栏")
st.markdown(
    f"**请在左侧手写签名，并在右侧手写日期（注：当前系统日期为"
    f" {current_date_str}，请按此手写日期）：**"
)

col_sig, col_date = st.columns(2)
with col_sig:
  st.markdown("**手写签名：**")
  canvas_result = st_canvas(
      stroke_width=4,
      stroke_color="#000000",
      background_color="#F8F9FA",
      height=200,
      width=320,
      drawing_mode="freedraw",
      key="canvas_sig",
      return_image_data=True,
  )
with col_date:
  st.markdown("**手写日期栏（请手写当前日期）：**")
  canvas_date_result = st_canvas(
      stroke_width=3,
      stroke_color="#000000",
      background_color="#F8F9FA",
      height=200,
      width=320,
      drawing_mode="freedraw",
      key="canvas_date",
      return_image_data=True,
  )


# ================= 7. 辅助函数：生成 Word 点检自检表 =================
def generate_self_monitoring_docx(
    person,
    emp_id,
    tasks,
    items,
    checks,
    sig_image_io,
    date_image_io,
):
  doc = Document()
  p_title = doc.add_paragraph()
  r_t = p_title.add_run("【自行监测与采样人员现场点检自检确认表】")
  r_t.bold = True
  r_t.font.size = Pt(16)

  for p in doc.paragraphs:
    for r in p.runs:
      r.font.name = "华文宋体"
      r.font.element.rPr.rFonts.set(qn("w:eastAsia"), "华文宋体")

  tasks_str = "、".join(tasks) if tasks else "未选择"

  check_details = "\n".join(
      [
          f"• {items[i]}: {'[符合]' if checks[i] else '[未勾选]'}"
          for i in range(len(items))
      ]
  )

  p_info = doc.add_paragraph()
  run_i = p_info.add_run(
      f"采样人员: {person}    工号/身份证: {emp_id}\n"
      f"监测任务/点位: {tasks_str}\n"
      f"填报时间: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
      "--------------------------------------------------\n"
      f"【现场检查要求落实情况】\n{check_details}\n\n"
      "本人承诺以上现场点检项目真实有效，严格遵循环保及IWAY合规标准。"
  )
  run_i.font.name = "华文宋体"
  run_i.font.size = Pt(10.5)
  run_i.font.element.rPr.rFonts.set(qn("w:eastAsia"), "华文宋体")

  p_line = doc.add_paragraph(
      "--------------------------------------------------"
  )
  p_line.paragraph_format.space_before = Pt(5)
  p_line.paragraph_format.space_after = Pt(10)

  table = doc.add_table(rows=1, cols=2)
  table.autofit = False

  cell_sig = table.cell(0, 0)
  p1 = cell_sig.paragraphs[0]
  r1 = p1.add_run("采样人员手写亲笔签名：\n")
  r1.font.name = "华文宋体"
  r1.font.size = Pt(10)
  r1.font.element.rPr.rFonts.set(qn("w:eastAsia"), "华文宋体")
  p1.add_run().add_picture(sig_image_io, width=Inches(1.8))
  sig_image_io.seek(0)

  cell_date = table.cell(0, 1)
  p2 = cell_date.paragraphs[0]
  r2 = p2.add_run("手写签署日期：\n")
  r2.font.name = "华文宋体"
  r2.font.size = Pt(10)
  r2.font.element.rPr.rFonts.set(qn("w:eastAsia"), "华文宋体")
  p2.add_run().add_picture(date_image_io, width=Inches(1.8))
  date_image_io.seek(0)

  buffer = io.BytesIO()
  doc.save(buffer)
  buffer.seek(0)
  return buffer


# ================= 8. 提交校验与生成档案 =================
if st.button("📁 确认无误，一键提交点检自检表并生成合规档案", use_container_width=True):
  is_canvas_empty = canvas_result.image_data is None or (
      canvas_result.json_data is not None
      and len(canvas_result.json_data.get("objects", [])) == 0
  )
  is_date_empty = canvas_date_result.image_data is None or (
      canvas_date_result.json_data is not None
      and len(canvas_date_result.json_data.get("objects", [])) == 0
  )

  if not sampling_person.strip() or not employee_id.strip() or not selected_tasks:
    st.error("❌ 拦截：请完整填写【采样人员姓名】、【工号/身份证】并【至少勾选一项监测任务】！")
  elif not all(user_checks):
    st.warning("⚠️ 拦截：为保障监测数据真实有效，必须将上方所有现场点检项全部勾选确认！")
  elif is_canvas_empty:
    st.warning("⚠️ 拦截：请在左侧画板完成手写签名后再提交。")
  elif is_date_empty:
    st.warning("⚠️ 拦截：请在右侧手写日期栏内完成手写日期后再提交！")
  else:
    st.success("✅ 点检自检表提交成功！系统已成功生成您的专属带签名 Word 合规确认档案。")

    signature_img = Image.fromarray(
        canvas_result.image_data.astype("uint8"), "RGBA"
    )
    sig_io = io.BytesIO()
    signature_img.save(sig_io, format="PNG")
    sig_io.seek(0)

    date_img = Image.fromarray(
        canvas_date_result.image_data.astype("uint8"), "RGBA"
    )
    date_io = io.BytesIO()
    date_img.save(date_io, format="PNG")
    date_io.seek(0)

    receipt_docx_buffer = generate_self_monitoring_docx(
        sampling_person,
        employee_id,
        selected_tasks,
        current_dynamic_items,
        user_checks,
        sig_io,
        date_io,
    )

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
      task_summary = "_".join(selected_tasks)[:15]
      receipt_filename = f"自行监测点检表_{sampling_person}_{task_summary}.docx"
      zip_file.writestr(receipt_filename, receipt_docx_buffer.getvalue())

      # 自动同步到百度网盘
      upload_to_baidu_netdisk_with_auto_refresh(
          receipt_docx_buffer.getvalue(), receipt_filename
      )

      img_byte_arr = io.BytesIO()
      signature_img.save(img_byte_arr, format="PNG")
      zip_file.writestr(
          f"手写签名原图_{sampling_person}.png", img_byte_arr.getvalue()
      )

      date_byte_arr = io.BytesIO()
      date_img.save(date_byte_arr, format="PNG")
      zip_file.writestr(f"手写日期原图_{sampling_person}.png", date_byte_arr.getvalue())

    zip_buffer.seek(0)

    st.markdown("---")
    st.success("🎉 您的自行监测点检档案已打包完毕，点击下方按钮即可下载保存！")

    col_d1, col_d2 = st.columns(2)
    with col_d1:
      st.download_button(
          label="📄 下载点检确认表 (.docx)",
          data=receipt_docx_buffer.getvalue(),
          file_name=f"自行监测点检表_{sampling_person}.docx",
          mime=(
              "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
          ),
          use_container_width=True,
      )
    with col_d2:
      st.download_button(
          label="📥 一键打包下载全部档案 (.ZIP)",
          data=zip_buffer,
          file_name=f"自行监测档案_{sampling_person}.zip",
          mime="application/zip",
          use_container_width=True,
      )

    st.snow()

# ================= 9. 底部版权与开发者声明 =================
st.markdown("---")
st.markdown(
    "<div style='text-align: center; color: gray; font-size: 14px;'>"
    "本系统为内部合规平台，严禁商业用途 | 开发者：陈野菲"
    "</div>",
    unsafe_allow_html=True,
)
